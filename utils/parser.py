"""
parser.py — Resume text extraction module.
Supports PDF, DOCX, and TXT formats.
"""

import pdfplumber
import docx
import io


def extract_text_from_pdf(file) -> str:
    """Extract text from a PDF file using pdfplumber."""
    text = ""
    try:
        with pdfplumber.open(file) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        raise ValueError(f"Error reading PDF: {e}")
    return text.strip()


def extract_text_from_docx(file) -> str:
    """Extract text from a DOCX file using python-docx."""
    text = ""
    try:
        # python-docx needs a file-like object or path
        doc = docx.Document(io.BytesIO(file.read()))
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text += cell.text + " "
                text += "\n"
    except Exception as e:
        raise ValueError(f"Error reading DOCX: {e}")
    return text.strip()


def extract_text_from_txt(file) -> str:
    """Extract text from a plain TXT file."""
    try:
        content = file.read()
        if isinstance(content, bytes):
            content = content.decode("utf-8", errors="ignore")
        return content.strip()
    except Exception as e:
        raise ValueError(f"Error reading TXT: {e}")


def extract_resume_text(uploaded_file) -> str:
    """
    Dispatch text extraction based on file extension.

    Args:
        uploaded_file: Streamlit UploadedFile object

    Returns:
        Extracted plain text string
    """
    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        return extract_text_from_pdf(uploaded_file)
    elif filename.endswith(".docx"):
        return extract_text_from_docx(uploaded_file)
    elif filename.endswith(".txt"):
        return extract_text_from_txt(uploaded_file)
    else:
        raise ValueError(f"Unsupported file format: {uploaded_file.name}. Use PDF, DOCX, or TXT.")
